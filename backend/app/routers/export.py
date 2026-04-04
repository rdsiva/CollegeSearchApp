import csv
import io
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.models.schemas import ExportRequest

router = APIRouter()

TNEA_CATEGORIES = ["OC", "BC", "BCM", "MBC", "SC", "ST", "SCA"]


def _get_cutoff(college, branch_code: str, category: str) -> str:
    for course in college.courses:
        if course.branch_code == branch_code:
            val = getattr(course.cutoffs, category, None)
            return str(val) if val is not None else "-"
    return "-"


@router.post("/export/csv")
async def export_csv(req: ExportRequest):
    """Export favorites as CSV in TNEA counselling template format."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Header row
    header = [
        "Institution Code", "Institution Name", "District", "City",
        "Affiliation", "NIRF Rank", "Branch Code", "Branch Name",
        "OC", "BC", "BCM", "MBC", "SC", "ST", "SCA",
        "TNEA Fees (₹)", "Management Fees (₹)", "Hostel Fees (₹)",
        "Avg LPA", "Highest LPA", "Top Companies",
        "Google Rating", "Sentiment Score", "College Score",
    ]
    writer.writerow(header)

    for college in req.colleges:
        branches = college.courses if college.courses else []
        if not branches:
            # Write one row without branch data
            row = [
                college.anna_university_code, college.name, college.district, college.city,
                college.affiliation, college.nirf_rank or "-",
                "-", "-",
                "-", "-", "-", "-", "-", "-", "-",
                college.fees.tnea if college.fees else "-",
                college.fees.management if college.fees else "-",
                college.fees.hostel if college.fees else "-",
                college.placement.avg_lpa if college.placement else "-",
                college.placement.highest_lpa if college.placement else "-",
                ", ".join(college.placement.top_companies) if college.placement else "-",
                college.reviews.google_rating if college.reviews else "-",
                college.reviews.sentiment_score if college.reviews else "-",
                college.score or "-",
            ]
            writer.writerow(row)
        else:
            for branch in branches:
                cutoffs = branch.cutoffs
                row = [
                    college.anna_university_code, college.name, college.district, college.city,
                    college.affiliation, college.nirf_rank or "-",
                    branch.branch_code, branch.branch_name,
                    cutoffs.OC or "-", cutoffs.BC or "-", cutoffs.BCM or "-",
                    cutoffs.MBC or "-", cutoffs.SC or "-", cutoffs.ST or "-", cutoffs.SCA or "-",
                    college.fees.tnea if college.fees else "-",
                    college.fees.management if college.fees else "-",
                    college.fees.hostel if college.fees else "-",
                    college.placement.avg_lpa if college.placement else "-",
                    college.placement.highest_lpa if college.placement else "-",
                    ", ".join(college.placement.top_companies) if college.placement else "-",
                    college.reviews.google_rating if college.reviews else "-",
                    college.reviews.sentiment_score if college.reviews else "-",
                    college.score or "-",
                ]
                writer.writerow(row)

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=college_research.csv"},
    )


def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    # Header row
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        run = hdr[i].paragraphs[0].runs[0] if hdr[i].paragraphs[0].runs else hdr[i].paragraphs[0].add_run(h)
        run.bold = True
    # Data rows
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = str(val)
    return table


@router.post("/export/word")
async def export_word(req: ExportRequest):
    """Export favorites as a formatted Word document."""
    doc = Document()
    doc.core_properties.title = "College Research Report"

    title = doc.add_heading("College Research Report — Tamil Nadu Engineering Colleges", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Summary table on page 1
    doc.add_heading("Summary", level=1)
    summary_headers = [
        "#", "Institution Code", "College Name", "District",
        "NIRF Rank", "TNEA Fees", "Avg LPA", "Sentiment", "Score"
    ]
    summary_rows = []
    for i, c in enumerate(req.colleges, 1):
        summary_rows.append([
            str(i),
            c.anna_university_code,
            c.name,
            c.district,
            str(c.nirf_rank or "-"),
            f"₹{c.fees.tnea:,}" if c.fees and c.fees.tnea else "-",
            f"{c.placement.avg_lpa} LPA" if c.placement and c.placement.avg_lpa else "-",
            f"{c.reviews.sentiment_score}/10" if c.reviews and c.reviews.sentiment_score else "-",
            f"{c.score}/100" if c.score else "-",
        ])
    _add_table(doc, summary_headers, summary_rows)

    # Detailed section per college
    for college in req.colleges:
        doc.add_page_break()
        _add_heading(doc, f"{college.name}", level=1)

        # Basic info
        doc.add_paragraph(f"Code: {college.anna_university_code}  |  Location: {college.city}, {college.district}")
        doc.add_paragraph(f"Affiliation: {college.affiliation}  |  NIRF Rank: {college.nirf_rank or 'Not ranked'}")
        doc.add_paragraph(f"Approved by: {', '.join(college.approved_by)}")

        # Score breakdown
        if college.score_breakdown:
            _add_heading(doc, "College Score", level=2)
            sb = college.score_breakdown
            score_headers = ["Component", "Weight", "Score"]
            score_rows = [
                ["Placement", "40%", f"{sb.placement}"],
                ["Fees Affordability", "20%", f"{sb.fees_affordability}"],
                ["Reviews Sentiment", "20%", f"{sb.reviews_sentiment}"],
                ["Infrastructure/Faculty", "10%", f"{sb.infrastructure}"],
                ["Location/Exposure", "10%", f"{sb.location}"],
                ["TOTAL", "100%", f"{sb.total}"],
            ]
            _add_table(doc, score_headers, score_rows)

        # Fees
        if college.fees:
            _add_heading(doc, "Fees", level=2)
            fee_headers = ["Quota", "Annual Fees"]
            fee_rows = []
            if college.fees.tnea:
                fee_rows.append(["TNEA (Counselling)", f"₹{college.fees.tnea:,}"])
            if college.fees.management:
                fee_rows.append(["Management Quota", f"₹{college.fees.management:,}"])
            if college.fees.hostel:
                fee_rows.append(["Hostel", f"₹{college.fees.hostel:,}"])
            if college.fees.transport:
                fee_rows.append(["Transport", f"₹{college.fees.transport:,}"])
            if fee_rows:
                _add_table(doc, fee_headers, fee_rows)

        # Placements
        if college.placement:
            _add_heading(doc, "Placement Statistics", level=2)
            pl = college.placement
            pl_headers = ["Metric", "Value"]
            pl_rows = []
            if pl.avg_lpa:
                pl_rows.append(["Average Package", f"₹{pl.avg_lpa} LPA"])
            if pl.highest_lpa:
                pl_rows.append(["Highest Package", f"₹{pl.highest_lpa} LPA"])
            if pl.placement_percentage:
                pl_rows.append(["Placement %", f"{pl.placement_percentage}%"])
            if pl.top_companies:
                pl_rows.append(["Top Recruiters", ", ".join(pl.top_companies)])
            if pl_rows:
                _add_table(doc, pl_headers, pl_rows)

        # Cutoffs
        if college.courses:
            _add_heading(doc, "Last Year TNEA Cutoffs", level=2)
            cutoff_headers = ["Branch", "OC", "BC", "BCM", "MBC", "SC", "ST", "SCA"]
            cutoff_rows = []
            for branch in college.courses:
                co = branch.cutoffs
                cutoff_rows.append([
                    branch.branch_name,
                    str(co.OC or "-"), str(co.BC or "-"), str(co.BCM or "-"),
                    str(co.MBC or "-"), str(co.SC or "-"), str(co.ST or "-"),
                    str(co.SCA or "-"),
                ])
            _add_table(doc, cutoff_headers, cutoff_rows)

        # Reviews
        if college.reviews:
            _add_heading(doc, "Student Reviews Summary", level=2)
            rv = college.reviews
            if rv.google_rating:
                doc.add_paragraph(f"Google Rating: {rv.google_rating}/5")
            if rv.sentiment_score is not None:
                doc.add_paragraph(f"Overall Sentiment Score: {rv.sentiment_score}/10")
            if rv.summary:
                doc.add_paragraph(rv.summary)
            if rv.pros:
                doc.add_paragraph("Pros:", style="Heading 3")
                for p in rv.pros:
                    doc.add_paragraph(p, style="List Bullet")
            if rv.cons:
                doc.add_paragraph("Cons:", style="Heading 3")
                for c in rv.cons:
                    doc.add_paragraph(c, style="List Bullet")
            if rv.common_complaints:
                doc.add_paragraph("Common Complaints:", style="Heading 3")
                for cc in rv.common_complaints:
                    doc.add_paragraph(cc, style="List Bullet")

        # YouTube videos
        if college.youtube_videos:
            _add_heading(doc, "Related YouTube Reviews", level=2)
            for v in college.youtube_videos:
                doc.add_paragraph(f"• {v.title}: {v.url}")

    # Serialize to bytes
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=college_research.docx"},
    )
